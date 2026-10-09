"""Project risk-assessment exports from the exact approved workflow snapshot."""

from html import escape
from io import BytesIO

from .workflow import WorkflowError


def report_tables(document):
    if document["status"] not in ("approved", "closed") or not document["approval"]:
        raise WorkflowError("관리감독자 승인 후 평가서를 출력할 수 있습니다.")
    form, approval = document["form"], document["approval"]
    source = document["source"]
    summary = [
        ["평가서 ID", document["id"]], ["문서 버전", document["version"]],
        ["상태", "종결" if document["status"] == "closed" else "승인 / 조치 진행"],
        ["사업장", form["site"]], ["공정", form["process"]], ["작업", form["task"]],
        ["평가일", form["assessment_date"]], ["참여자", form["participants"]],
        ["평가 방법", "프로젝트 기본 5×5: 가능성(1–5) × 중대성(1–5), 사람이 입력"],
        ["허용 판단 기준", form["acceptance_criteria"]], ["현장 확인 의견", form["review_note"]],
        ["관리감독자", approval["actor"]], ["승인 시각 (UTC)", approval["timestamp"]],
        ["승인 대상 버전", approval["reviewed_version"]], ["승인 의견", approval["note"]],
        ["관리감독자 확인 항목", "원본 현장자료, 위험요인 누락·오탐, 개선조치, 위험도·허용 기준 확인"],
        ["입력 파일", source.get("input_name", "")], ["입력 SHA-256", source.get("input_sha256", "")],
        ["모델", source.get("model_id", "")], ["LoRA 적용", str(source.get("adapter_loaded", False))],
        ["분석 감사 기록", str(source.get("audit") or "미설정")],
        ["양식 버전", document["template_version"]],
        ["변경 이력 최종 해시", document["history"][-1]["hash"]],
        ["양식 안내", "프로젝트 기본 양식. 회사 지정 양식 및 평가 기준은 별도 확인."],
    ]
    if document.get("closure"):
        summary.extend([["종결 확인자", document["closure"]["actor"]],
                        ["종결 시각 (UTC)", document["closure"]["timestamp"]],
                        ["종결 의견", document["closure"]["note"]]])
    risks = [["항목 ID", "위험요인", "예상 피해", "현재 안전조치", "가능성", "중대성", "위험도",
              "추가 개선조치", "조치 담당자", "기한", "조치 상태", "완료 내용", "완료 증빙",
              "완료 등록자", "완료 시각 (UTC)", "잔여 가능성", "잔여 중대성", "잔여 위험도",
              "재확인자", "재확인 시각 (UTC)", "재확인 의견"]]
    labels = {"open": "미완료", "completed": "재확인 대기", "verified": "재확인 완료"}
    for row in form["rows"]:
        action = document["actions"][row["id"]]
        residual = (action["residual_likelihood"] * action["residual_severity"]
                    if "residual_likelihood" in action else "")
        risks.append([
            row["id"], row["hazard"], row["consequence"], row["existing_controls"],
            row["likelihood"], row["severity"], row["likelihood"] * row["severity"],
            row["additional_controls"], row["owner"], row["due_date"], labels[action["status"]],
            action.get("note", ""), action.get("evidence", ""), action.get("completed_by", ""),
            action.get("completed_at", ""), action.get("residual_likelihood", ""),
            action.get("residual_severity", ""), residual, action.get("verified_by", ""),
            action.get("verified_at", ""), action.get("verification_note", ""),
        ])
    history = [["순번", "시각 (UTC)", "사용자", "역할", "작업", "의견", "해시"]]
    history.extend([[e["seq"], e["timestamp"], e["actor"]["id"], e["actor"]["role"],
                     e["type"], e["note"], e["hash"]] for e in document["history"]])
    return [("평가 개요", summary), ("위험성평가 및 개선조치", risks), ("확인 이력", history)]


def export_xlsx(document):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    for title, rows in report_tables(document):
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                # Treat user/model text as literal text, never a spreadsheet formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 28 if i > 1 else 22
        sheet.freeze_panes = "B2"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:1"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def export_html(document):
    sections = []
    for title, rows in report_tables(document):
        if title == "위험성평가 및 개선조치":
            # Cards keep every field readable when printing on A4 paper.
            content = "".join("<article><h3>위험요인 " + str(i) + "</h3><table>" + "".join(
                f"<tr><th>{escape(str(k))}</th><td>{escape(str(v))}</td></tr>"
                for k, v in zip(rows[0], row)
            ) + "</table></article>" for i, row in enumerate(rows[1:], 1))
            if len(rows) == 1:
                content = "<p>등록된 위험요인 없음. 현장 확인 및 승인 의견을 참조하세요.</p>"
        else:
            content = "<table>" + "".join("<tr>" + "".join(
                f"<td>{escape(str(cell))}</td>" for cell in row
            ) + "</tr>" for row in rows) + "</table>"
        sections.append(f"<section><h2>{escape(title)}</h2>{content}</section>")
    return """<!doctype html><html lang="ko"><meta charset="utf-8">
<title>위험성평가서</title><style>
body{font-family:system-ui,sans-serif;max-width:1100px;margin:32px auto;color:#172033}
table{border-collapse:collapse;width:100%;font-size:12px;table-layout:fixed}
td,th{border:1px solid #b8c3d3;padding:8px;text-align:left;white-space:pre-wrap;overflow-wrap:anywhere}
th{width:22%;background:#eef2ff}article{margin:20px 0}h2{color:#1d4ed8}
@media print{button{display:none}body{margin:0}tr{break-inside:avoid}h2,h3{break-after:avoid}}
@page{size:A4;margin:12mm}
</style><button onclick="window.print()">인쇄 / PDF로 저장</button>
<h1>위험성평가서</h1><p>관리감독자 승인본 · 프로젝트 기본 양식</p>""" + "".join(sections) + "</html>"
