import copy
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.workflow import Conflict, WorkflowStore
from server import create_app

TOKENS = {'author': 'a' * 32, 'boss': 'b' * 32, 'other': 'c' * 32}
USERS = {token: {'id': name, 'role': 'operator' if name == 'author' else 'supervisor'} for name, token in TOKENS.items()}
RESULT = {'hazards_detected': [{'type': '끼임', 'event_type': 'entrapment', 'description': '회전부 접근', 'severity': 'High'}],
          'overall_safety_score': 30, 'recommendation': '방호장치 설치'}


class InferenceStub:
    model = object()
    model_id = 'test-inference-stub'
    adapter_loaded = False
    use_label_grounding = False
    result = RESULT

    def analyze_media(self, path, safety_context):
        yield {'stage': 'preprocessing', 'detail': 'test'}
        yield {'stage': 'complete', 'result': copy.deepcopy(self.result)}


def auth(name='author'):
    return {'Authorization': f'Bearer {TOKENS[name]}'}


@pytest.fixture
def runtime(tmp_path):
    agent = InferenceStub()
    app = create_app(agent=agent, db_path=tmp_path / 'workflow.db', users=USERS,
                     upload_dir=tmp_path / 'uploads', inference_enabled=True)
    with TestClient(app) as client:
        yield client, app.state.store, agent


def create(client):
    response = client.post('/analyze', headers=auth(), files={'file': ('../frame.png', b'pixels', 'image/png')})
    assert response.status_code == 200
    events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: ')]
    assert events[-1]['stage'] == 'complete', events
    return client.get('/assessments/' + events[-1]['assessment_id'], headers=auth()).json()


def filled(document):
    form = copy.deepcopy(document['form'])
    form.update(site='현장 A', process='프레스', task='정비', assessment_date='2026-10-09',
                participants='작업자, 관리감독자', supervisor_id='boss',
                acceptance_criteria='현장 자체 척도와 허용 기준', review_note='원본과 현장 상황 확인')
    for row in form['rows']:
        row.update(consequence='손 부상', existing_controls='접근 제한', likelihood=3, severity=4,
                   additional_controls='방호장치 설치', owner='author', due_date='2026-10-10')
    return form


def mutate(client, document, path, name='author', **payload):
    if path == '/decisions/approve' and 'checks' not in payload:
        payload['checks'] = {key: True for key in ('source_checked', 'hazards_checked', 'controls_checked', 'criteria_checked')}
    return client.post(f"/assessments/{document['id']}{path}", headers=auth(name),
                       json={'version': document['version'], **payload})


def pending(client):
    document = create(client)
    response = client.put('/assessments/' + document['id'], headers=auth(),
                          json={'version': document['version'], 'form': filled(document)})
    assert response.status_code == 200, response.text
    response = mutate(client, response.json(), '/submit')
    assert response.status_code == 200, response.text
    return response.json()


def approved(client):
    response = mutate(client, pending(client), '/decisions/approve', name='boss', note='현장 확인 후 승인')
    assert response.status_code == 200, response.text
    return response.json()


def test_analysis_persists_provenance_and_survives_restart(runtime):
    client, store, _ = runtime
    document = create(client)
    assert document['status'] == 'draft'
    assert document['original_result'] == RESULT
    assert document['source']['input_name'] == 'frame.png'
    assert document['source']['input_sha256'] == hashlib.sha256(b'pixels').hexdigest()
    assert document['form']['rows'][0]['likelihood'] is None
    assert len(list((store.path.parent / 'uploads').iterdir())) == 1
    assert client.get(f"/assessments/{document['id']}/media", headers=auth('boss')).content == b'pixels'
    assert client.get(f"/assessments/{document['id']}/media").status_code == 401
    assert WorkflowStore(store.path).get(document['id']) == document


def test_full_supervised_workflow_exports_and_closes(runtime):
    client, store, _ = runtime
    document = approved(client)
    original_form = copy.deepcopy(document['form'])
    assert document['approval']['actor'] == 'boss'
    assert mutate(client, document, '/decisions/close', name='boss', note='종결').status_code == 409
    row_id = document['form']['rows'][0]['id']
    document = mutate(client, document, f'/actions/{row_id}/complete', note='설치 완료', evidence='현장 사진 IMG_1',
                      residual_likelihood=1, residual_severity=4).json()
    assert document['actions'][row_id]['status'] == 'completed'
    assert mutate(client, document, f'/actions/{row_id}/verify', note='확인').status_code == 403
    document = mutate(client, document, f'/actions/{row_id}/verify', name='boss', note='잔여 위험 허용 확인').json()
    document = mutate(client, document, '/decisions/close', name='boss', note='모든 조치 확인').json()
    assert document['status'] == 'closed'
    assert document['form'] == original_form
    assert store.verify(document['id'])
    xlsx = client.get(f"/assessments/{document['id']}/export/xlsx", headers=auth())
    assert xlsx.status_code == 200
    book = load_workbook(BytesIO(xlsx.content))
    assert book['위험성평가 및 개선조치']['G2'].value == 12
    assert book['위험성평가 및 개선조치']['R2'].value == 4
    assert book['위험성평가 및 개선조치']['S2'].value == 'boss'
    html = client.get(f"/assessments/{document['id']}/export/html", headers=auth())
    assert html.status_code == 200 and '인쇄 / PDF로 저장' in html.text and '모든 조치 확인' in html.text
    assert client.put('/assessments/' + document['id'], headers=auth(), json={
        'version': document['version'], 'form': original_form}).status_code == 409


@pytest.mark.parametrize('name', ['author', 'other'])
def test_only_assigned_supervisor_can_approve(runtime, name):
    client, _, _ = runtime
    document = pending(client)
    assert mutate(client, document, '/decisions/approve', name=name, note='승인').status_code == 403


def test_reject_revision_resubmit_and_stale_approval(runtime):
    client, _, _ = runtime
    document = pending(client)
    stale = copy.deepcopy(document)
    document = mutate(client, document, '/decisions/request_changes', name='boss', note='조치 보완').json()
    assert document['status'] == 'changes_requested'
    document = mutate(client, document, '/submit').json()
    assert mutate(client, stale, '/decisions/approve', name='boss', note='이전 버전 승인').status_code == 409
    assert mutate(client, document, '/decisions/approve', name='boss', note='재검토 승인').status_code == 200


def test_required_fields_and_unapproved_export(runtime):
    client, _, _ = runtime
    document = create(client)
    assert mutate(client, document, '/submit').status_code == 422
    assert client.get(f"/assessments/{document['id']}/export/xlsx", headers=auth()).status_code == 422
    form = filled(document)
    form['rows'][0]['likelihood'] = 6
    assert client.put('/assessments/' + document['id'], headers=auth(), json={
        'version': document['version'], 'form': form}).status_code == 422


@pytest.mark.parametrize('result', [{'error': 'bad'}, {}, {'hazards_detected': [{}]}, {'hazards_detected': 'missing'}])
def test_failed_analysis_does_not_create_safe_record(runtime, result):
    client, store, agent = runtime
    agent.result = result
    response = client.post('/analyze', headers=auth(), files={'file': ('frame.png', b'pixels')})
    assert '"stage": "error"' in response.text
    assert 'assessment_id' not in response.text
    assert store.list() == []


def test_authentication_identity_spoof_and_disabled_inference(runtime, tmp_path):
    client, _, _ = runtime
    assert client.get('/assessments').status_code == 401
    document = pending(client)
    assert mutate(client, document, '/decisions/approve', note='승인', actor='boss').status_code == 422
    app = create_app(users=USERS, db_path=tmp_path / 'review.db', inference_enabled=False)
    with TestClient(app) as review:
        assert review.post('/analyze', headers=auth(), files={'file': ('a.png', b'a')}).status_code == 503
    with TestClient(create_app(users={}, db_path=tmp_path / 'unconfigured.db')) as empty:
        assert empty.get('/assessments').status_code == 503


def test_reopen_requires_new_completion(runtime):
    client, _, _ = runtime
    document = approved(client)
    row_id = document['form']['rows'][0]['id']
    document = mutate(client, document, f'/actions/{row_id}/complete', note='완료', evidence='사진',
                      residual_likelihood=2, residual_severity=4).json()
    document = mutate(client, document, f'/actions/{row_id}/reopen', name='boss', note='추가 보완').json()
    assert document['actions'][row_id] == {'status': 'open', 'reopen_note': '추가 보완'}
    assert mutate(client, document, f'/actions/{row_id}/verify', name='boss', note='확인').status_code == 409


def test_concurrent_edits_cannot_overwrite_review(runtime):
    client, store, _ = runtime
    document = create(client)
    def edit():
        try:
            store.mutate(document['id'], document['version'], USERS[TOKENS['author']], 'edit', form=filled(document))
            return 'ok'
        except Conflict:
            return 'conflict'
    with ThreadPoolExecutor(max_workers=2) as pool:
        outcomes = list(pool.map(lambda _: edit(), range(2)))
    assert sorted(outcomes) == ['conflict', 'ok']
    assert store.verify(document['id'])


def test_tampering_blocks_export_and_mutations(runtime):
    client, store, _ = runtime
    document = approved(client)
    with store.connect() as db:
        db.execute("UPDATE events SET body='{}' WHERE assessment_id=? AND seq=0", (document['id'],))
    assert not store.verify(document['id'])
    assert client.get(f"/assessments/{document['id']}/export/html", headers=auth()).status_code == 409
    assert mutate(client, document, '/decisions/close', name='boss', note='확인').status_code == 409


def test_exports_escape_html_and_spreadsheet_formulas(runtime):
    client, _, agent = runtime
    agent.result = {'hazards_detected': [], 'recommendation': ''}
    document = create(client)
    form = filled(document)
    form.update(site='=HYPERLINK("https://example.com")', review_note='<script>alert(1)</script>')
    document = client.put('/assessments/' + document['id'], headers=auth(), json={
        'version': document['version'], 'form': form}).json()
    document = mutate(client, document, '/submit').json()
    document = mutate(client, document, '/decisions/approve', name='boss', note='위험 미검출 현장 확인').json()
    book = load_workbook(BytesIO(client.get(f"/assessments/{document['id']}/export/xlsx", headers=auth()).content))
    assert book['평가 개요']['B4'].data_type == 's'
    html = client.get(f"/assessments/{document['id']}/export/html", headers=auth()).text
    assert '<script>alert(1)</script>' not in html
    assert '&lt;script&gt;' in html
    assert mutate(client, document, '/decisions/close', name='boss', note='위험 없음 재확인').status_code == 200


def test_supervisor_must_complete_review_checklist(runtime):
    client, _, _ = runtime
    document = pending(client)
    assert mutate(client, document, '/decisions/approve', name='boss', note='확인', checks=None).status_code == 422
    checks = {'source_checked': True, 'hazards_checked': False, 'controls_checked': True, 'criteria_checked': True}
    assert mutate(client, document, '/decisions/approve', name='boss', note='확인', checks=checks).status_code == 422


def test_changed_source_media_is_not_served(runtime):
    client, store, _ = runtime
    document = create(client)
    path = store.path.parent / 'uploads' / document['source']['media_file']
    path.write_bytes(b'changed')
    assert client.get(f"/assessments/{document['id']}/media", headers=auth()).status_code == 409
